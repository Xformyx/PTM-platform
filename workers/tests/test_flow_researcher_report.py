"""Researcher report regressions: actual helpers, synthetic measured evidence."""
import ast
import asyncio
import logging
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
from common.species_detector import detect_species_from_tsv
from common.section_budgets import section_content_issues
from report_generation.core.scientific_semantics import _compress_question_answers


def test_gene_capitalization_is_not_an_annotation_species():
    for genes in (["Mapk1", "Mapk3", "Larp1"], ["MAPK1", "MAPK3", "LARP1"]):
        assert detect_species_from_tsv(pd.DataFrame({"Gene.Name": genes}))[0] == "unknown"


def uniprot_parse(comments):
    source = Path(__file__).parents[2] / "mcp-server/app/tools/uniprot.py"
    tree = ast.parse(source.read_text())
    nodes = [n for n in tree.body if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))
             and n.name in {"_fetch_uniprot_info", "_empty_result"}]
    class Client:
        def __init__(self, **kwargs): pass
        async def __aenter__(self): return self
        async def __aexit__(self, *args): pass
        async def get(self, url):
            return SimpleNamespace(status_code=200, headers={}, content=__import__("json").dumps(comments).encode(), raise_for_status=lambda: None,
                                   json=lambda: {"comments": comments, "organism": {"taxonId": 10116}})
    env = {"httpx": SimpleNamespace(AsyncClient=Client), "BASE_URL": "mock", "logger": logging.getLogger(__name__)}
    exec(compile(ast.Module(body=nodes, type_ignores=[]), str(source), "exec"), env)
    return asyncio.run(env["_fetch_uniprot_info"]("PTEST1", 1))


def test_uniprot_preserves_general_and_isoform_functions_in_either_order():
    comments = [{"commentType": "FUNCTION", "texts": [{"value": "General kinase function."}]},
                {"commentType": "FUNCTION", "molecule": "Isoform 2", "texts": [{"value": "Lacks the kinase domain."}]}]
    first, second = uniprot_parse(comments), uniprot_parse(comments[::-1])
    assert first["function_summary"] == second["function_summary"] == "General kinase function."
    assert len(first["function_comments"]) == 2
    assert first["function_comments"] == second["function_comments"]


def test_long_legacy_questions_never_erase_every_answer():
    answer = "Observed precursor contrasts differ between the recorded conditions and require matched measurements for independent validation."
    text = "\n\n".join(f"### Q{i}. " + "question context " * 18 + "\n\n" + answer for i in range(6))
    compressed = _compress_question_answers(text, 220)
    assert compressed.count(answer) == 6


def test_short_abstract_and_introduction_report_content_gap():
    assert "below_section_target" in section_content_issues("Measured features changed. " * 40, "abstract")
    assert "below_section_target" in section_content_issues("Study background context. " * 60, "introduction")


def test_supported_partial_early_pattern_survives_many_repeated_late_patterns():
    from report_generation.core.reader_authoring import build_authoring_packet, deterministic_authoring_plan
    state = researcher_state(50)
    state['research_questions'] = []
    for row in state['vector_plot_raw_data']:
        if row['gene'] != 'GENE0':
            value = {'1min': .2, '5min': .4, '15min': .6, '30min': .8, '60min': 1., '180min': 1.2}[row['condition']]
            row.update(ptm_unadjusted_log2fc=value, ptm_protein_adjusted_log2fc=value)
    packet = build_authoring_packet(state)
    plan = deterministic_authoring_plan(packet)
    selected = {f['reader_feature_id'] for f in plan['key_findings']}
    early = next(c for c in packet['reader_cards'] if c.get('feature_identity', {}).get('gene') == 'GENE0' and c.get('trajectory'))
    assert early['feature_identity']['reader_feature_id'] in selected
    assert next(p for p in early['trajectory'] if p['condition'] == '60min')['ptm_protein_adjusted_log2fc'] is None
    assert packet['observation_selection_audit']['input_unique_feature_count'] == 50


def researcher_state(count=4):
    from test_task10_joint_patterns import supported_row
    rows = []
    for i in range(count):
        for j, t in enumerate((1, 5, 15, 30, 60, 180)):
            a = None if t == 60 and i == 0 else ([.5, 1.7, 1.2, .6, .3, .2][j] if i % 2 == 0 else -.2 * (j + 1))
            rows.append(supported_row(precursor=f'form-{i}', gene=f'GENE{i}', condition=f'{t}min',
                        protein_group=f'PG{i}', protein_log2fc=.05,
                        ptm_unadjusted_log2fc=a, ptm_protein_adjusted_log2fc=a,
                        ptm_protein_adjusted_q_value=.001 if a is not None else None,
                        protein_method='ratio_of_condition_arithmetic_means; no_protein_test_computed',
                        protein_control_n=3, protein_treatment_n=3))
    return {'vector_plot_raw_data': rows, 'experimental_context': {'organism': 'rat', 'transgene_species': {'INSR': 'human'}, 'cell_type': 'synthetic cells'},
            'research_questions': ['How does protein adjustment change the observed GENE0 PTM contrasts?',
                                   'Does SRC, EGFR or PTK2 activation explain the response?']}


def test_explicit_native_species_preserves_host_transgene_and_mixed_groups():
    from ptm_shared.annotation_species import annotation_scope
    context = {'host_organism': 'rat', 'transgene_species': {'INSR': 'human'}}
    rat = annotation_scope({'FASTA_Taxonomy_ID': '10116'}, context)
    human = annotation_scope({'FASTA_Taxonomy_ID': '9606'}, context)
    mixed = annotation_scope({'FASTA_Taxonomy_ID': '10116;9606'}, context)
    assert rat['annotation_taxonomy_id'] == '10116'
    assert human['annotation_taxonomy_id'] == '9606' and human['host_taxonomy_id'] == '10116'
    assert mixed['annotation_taxonomy_id'] is None and mixed['status'] == 'mixed_or_ambiguous'
    assert annotation_scope({}, {})['native_species'] == 'unknown'
    assert rat['cache_namespace'] != human['cache_namespace'] != mixed['cache_namespace']
    assert detect_species_from_tsv(pd.DataFrame({'KEGG': ['signaling (mmu)']}))[0] == 'unknown'


def test_stale_mouse_annotation_cannot_be_silently_reused_for_rat_report():
    from report_generation.core.reader_authoring import build_authoring_packet
    state = researcher_state(1)
    state['parsed_ptms'] = [{'FASTA_Taxonomy_ID': '10116', 'rag_enrichment': {'native_species': 'mouse'}}]
    metadata = build_authoring_packet(state)['study_metadata_contract']
    assert metadata['native_annotation_audit']['mismatched_records'] == 1
    assert 'native_annotation_revalidation_required' in metadata['review_reason_codes']


def test_unrelated_findings_cannot_answer_a_gene_specific_question():
    from report_generation.core.reader_authoring import build_authoring_packet, render_reader_section_fallback
    from report_generation.core.research_questions import audit_question_coverage
    packet = build_authoring_packet(researcher_state())
    mapping = packet['research_question_evidence_map']
    assert mapping['questions'][0]['feature_ids']
    assert mapping['questions'][1]['feature_ids'] == []
    sections = {s: render_reader_section_fallback(s, packet) for s in ('results', 'discussion')}
    audit = audit_question_coverage(sections, mapping)
    assert audit['questions'][0]['answer_status'] == 'answered'
    assert audit['questions'][1]['answer_status'] == 'unanswered'
    assert audit['questions'][1]['discussion_paragraph_ids']
    wrong = audit_question_coverage({'results': sections['results'], 'discussion': 'GENE0 has changed in abundance.'}, mapping)
    assert wrong['questions'][1]['coverage_status'] == 'missing'
    assert 'Q1' not in sections['discussion'] and 'Supplementary Research Question Answers' not in sections['discussion']


def test_available_u_cannot_be_described_as_unavailable():
    from report_generation.core.reader_authoring import build_authoring_packet
    from report_generation.core.quantitative_claims import validate_quantitative_sentence
    packet = build_authoring_packet(researcher_state())
    assert 'available_unadjusted_observation_denied' in validate_quantitative_sentence('No independent unadjusted PTM comparison was available in this run.', packet)


def test_focused_structured_prompt_preserves_facts_without_duplicate_catalog():
    import json
    from report_generation.core.reader_authoring import build_authoring_packet, deterministic_authoring_plan, focus_authoring_packet, format_authoring_packet_for_llm
    from report_generation.core.quantitative_claims import structured_authoring_instructions, value_token_catalog, quantitative_records
    packet = build_authoring_packet(researcher_state(16))
    plan = deterministic_authoring_plan(packet)
    focused = focus_authoring_packet(packet, plan)
    catalog = value_token_catalog(focused)
    instructions = structured_authoring_instructions(focused)
    serialized, unavailable = instructions.split('Immutable token references: ', 1)[1].split('\nUnavailable observations (not value tokens; preserve the supplied reasons): ')
    assert json.loads(serialized) == catalog  # p/q/n, full sample support, unit and identity remain intact.
    assert json.loads(unavailable) == [r for c in focused['reader_cards'] for r in quantitative_records(c) if r['value'] is None]
    for section in ('results', 'discussion', 'conclusion', 'abstract'):
        prompt = format_authoring_packet_for_llm(focused, section, plan, include_quantitative_records=False) + instructions
        assert len(prompt) + 40_000 < 200_000  # Room for validated Results/Discussion and context.
        assert 'Quantitative references: ' not in prompt
    assert len(packet['reader_cards']) > len(focused['reader_cards'])  # Complete audit packet retained.


def test_empty_image_path_never_registers_output_directory(tmp_path):
    from report_generation.core.figure_manifest import _available_path
    assert _available_path('') == ''
    assert _available_path(str(tmp_path)) == ''


def test_heatmap_order_uses_elapsed_time_instead_of_lexical_labels():
    from report_generation.core.figure_manifest import order_reader_conditions
    labels = ['15min', '180min', '1min', '30min', '5min', '60min']
    assert order_reader_conditions(labels, []) == ['1min', '5min', '15min', '30min', '60min', '180min']
    assert order_reader_conditions(['late', 'early', 'unknown'], [{'condition': 'late', 'time_minutes': 40}, {'condition': 'early', 'time_minutes': 5}]) == ['early', 'late', 'unknown']


def test_study_search_is_cached_and_failed_comparison_keeps_provider_text():
    from report_generation.core.finding_literature import retrieve_finding_literature
    from report_generation.core.measured_feature_cards import build_feature_observation_cards
    calls = []
    class Retriever:
        collection_names = ['fixture']
        def query(self, query, **kwargs):
            calls.append(query)
            return [{'doi': '10.1000/fixture', 'document': 'GENE0 and GENE1 protein function in rat cells.', 'source_id': 'fixture'}]
    class Model:
        model, provider = 'fixture', 'mock'
        def generate(self, *args, **kwargs): return '{malformed'
    result = retrieve_finding_literature(build_feature_observation_cards(researcher_state(2)), Retriever(), {'organism': 'rat', 'cell_type': 'fixture cells'}, llm=Model())
    # One shared study-context query is cached; the other three retrieval
    # layers remain feature specific for each of the two named observations.
    assert len(calls) == 7 and sum('signaling time course' in c for c in calls) == 1
    assert 'rat' in calls[0] and all('mouse' not in c for c in calls)
    for record in result['records'].values():
        assert record['comparison_generation']['provider_raw_text'] == '{malformed'
        assert record['comparison_failure_type'] == 'JSONDecodeError'


def test_docx_keeps_figure_heading_and_caption_with_image_and_native_links(tmp_path):
    from PIL import Image
    from docx import Document
    from common.markdown_to_docx import convert_markdown_to_docx
    png, path = tmp_path / 'fixture.png', tmp_path / 'fixture.docx'
    Image.new('RGB', (600, 300), 'white').save(png)
    convert_markdown_to_docx(f'### Figure 1. Recorded contrasts\n\n![Recorded contrasts]({png})\n\n**Figure legend.** Observed measurements.\n\n[Source](https://example.org/source)', str(path))
    doc = Document(path)
    assert doc.paragraphs[0].paragraph_format.keep_with_next
    image_paragraph = next(p for p in doc.paragraphs if p._p.xpath('.//w:drawing'))
    assert image_paragraph.paragraph_format.keep_with_next
    assert len(doc.element.xpath('.//w:hyperlink')) == 1
    assert not any('[Source](' in p.text for p in doc.paragraphs)


def test_repaired_paragraph_extension_does_not_repeat_valid_prefix():
    import json
    from report_generation.core.quantitative_claims import merge_valid_sentence_drafts
    packet = {'reader_cards': [], 'figure_cards': []}
    one = {'text': 'The experiment compares recorded precursor contrasts.', 'scope': 'study_rationale', 'paragraph': 1, 'evidence_ids': [], 'value_tokens': [], 'figure_keys': []}
    two = {**one, 'text': 'The study asks how the measured response varies across conditions.'}
    prose, audit = merge_valid_sentence_drafts([json.dumps({'sentences': [one]}), json.dumps({'sentences': [one, two]})], packet)
    assert prose.count(one['text']) == 1 and two['text'] in prose
    assert len(audit) == 3


def test_provider_exception_malformed_json_and_valid_sentences_remain_distinct(monkeypatch):
    import json
    from common.llm_client import LLMClient
    from report_generation.core.reader_authoring import build_authoring_packet
    from report_generation.core.quantitative_claims import decode_sentence_draft, merge_valid_sentence_drafts
    packet = build_authoring_packet(researcher_state())
    sentence = {'text': 'The observed precursor contrast remained available for comparison across the measured times.', 'paragraph': 1,
                'scope': 'observation', 'evidence_ids': [], 'value_tokens': [], 'figure_keys': []}
    response = json.dumps({'sentences': [sentence, {**sentence, 'value_tokens': ['V999999']}]})
    outcomes = iter([TimeoutError('fixture'), '{broken', response])
    prompts = []
    client = LLMClient(provider='gemini', model='fixture-model', api_key='synthetic-only')
    def generate(**kwargs):
        prompts.append(kwargs['prompt'])
        result = next(outcomes)
        if isinstance(result, Exception): raise result
        return result
    monkeypatch.setattr(client, 'generate', generate)
    trace = []
    client.generate_with_retry('fixture', max_retries=3, trace_sink=trace,
        content_validator=lambda d: ['invalid_sentence'] if any(not r['retained'] for r in decode_sentence_draft(d, packet)[1]) else [])
    assert trace[0]['exception_type'] == 'TimeoutError'
    assert trace[1]['provider_raw_text'] == '{broken'
    assert trace[2]['provider_raw_text'] == response
    assert '{broken' in prompts[-1] and 'Prior draft to repair' in prompts[-1]
    prose, audit = merge_valid_sentence_drafts([r['provider_raw_text'] for r in trace], packet)
    assert prose.count(sentence['text']) == 1
    assert any(r.get('reason_code') == 'invalid_sentence_json' for r in audit)
    assert any(r.get('reason_codes') == ['unknown_or_unbound_value_token'] for r in audit)


def test_rejected_sibling_cannot_supply_hypothesis_paragraph_evidence():
    import json
    from report_generation.core.reader_authoring import build_authoring_packet
    from report_generation.core.quantitative_claims import decode_sentence_draft
    packet = build_authoring_packet(researcher_state())
    card = next(c for c in packet['reader_cards'] if c.get('trajectory'))
    packet['reader_cards'].append({'citation_ids': ['pmid:1'], 'evidence_ids': ['literature.pmid:1']})
    invalid = {'text': 'A claimed observation [REF:pmid:1].', 'paragraph': 1, 'scope': 'observation',
               'evidence_ids': card['evidence_ids'], 'value_tokens': ['V999999'], 'figure_keys': []}
    hypothesis = {**invalid, 'text': 'We predict a precursor-specific response that can be tested independently.', 'scope': 'testable_hypothesis', 'evidence_ids': [], 'value_tokens': []}
    prose, audit = decode_sentence_draft(json.dumps({'sentences': [invalid, hypothesis]}), packet)
    assert not prose and all(not r['retained'] for r in audit)


def test_protein_test_absence_is_distinct_from_missing_value_and_q_zero(tmp_path):
    from report_generation.core.figure_manifest import _generate_joint_trajectory_entry, joint_trajectory_evidence_table
    state = researcher_state(1)
    state['vector_plot_raw_data'][0]['ptm_unadjusted_q_value'] = 0
    figure = _generate_joint_trajectory_entry(state, str(tmp_path))
    rows = figure['quantitative_bindings']
    p = next(r for r in rows if r['axis'] == 'protein')
    assert p['value'] == .05 and p['support']['test_status'] == 'not_computed'
    assert p['support']['value_status'] == 'observed'
    assert next(r for r in rows if r['axis'] == 'unadjusted' and r['condition'] == '1min')['q'] == 0
    missing = next(r for r in rows if r['axis'] == 'unadjusted' and r['condition'] == '60min')
    assert missing['value'] is None
    table = joint_trajectory_evidence_table(figure)
    assert '| P: value; n |' in table and 'separate protein test was not computed' in table
    assert '; 0 |' in table


def test_finding_figures_share_nonempty_bindings_and_resolve_early_ticks(tmp_path):
    from report_generation.core.figure_manifest import prepare_reader_figure_manifest
    state = {**researcher_state(16), 'output_dir': str(tmp_path)}
    manifest = prepare_reader_figure_manifest(state, citation_complete=False)
    figures = {f['figure_key']: f for f in manifest['figures']}
    joint = figures['reader_joint_trajectories']
    assert joint['placement'] == 'main', joint
    assert joint['layout'] == 'overlaid_axes_with_early_inset'
    assert not joint['readability_audit']['tick_overlap_axes']
    for key in ('reader_quantitative_heatmap', 'reader_protein_context', 'reader_joint_trajectories'):
        assert figures[key]['quantitative_bindings']
    assert set(joint['selected_reader_feature_ids']).issubset(figures['reader_quantitative_heatmap']['selected_reader_feature_ids'])
    assert set(figures['reader_protein_context']['selected_reader_feature_ids']).issubset(joint['selected_reader_feature_ids'])
    lookup = {(r['feature_id'], r['condition'], r['axis']): r['value'] for r in joint['quantitative_bindings']}
    for r in figures['reader_protein_context']['quantitative_bindings']:
        assert lookup[r['feature_id'], r['condition'], r['axis']] == r['value']


def test_concordance_uses_integer_counts_and_preserves_all_intervals(tmp_path):
    from report_generation.core.figure_manifest import concordance_bindings, _generate_concordance_summary, FigureEligibilityPolicy
    rows = [{'static_wave_id': 'C1', 'from_window': a, 'to_window': b,
             'pair_transition_type_counts': {'persistence': 2, 'split': 1}, 'evaluable_pair_window_comparison_count': 7,
             'concordance_change_rates': {'retained': .99, 'gain': .99, 'loss': .99}}
            for a, b in [('15-30min', '30-60min'), ('1-5min', '5-15min')]]
    state = {'temporal_ptm_protein_analysis': {'dynamic_transition_per_wave': rows}}
    binding = concordance_bindings(state)
    assert binding[0]['from_window'] == '1-5min'
    assert binding[0]['retained'] == 2 and binding[0]['other_or_no_transition'] == 4
    path, clusters = _generate_concordance_summary(state, str(tmp_path), [])
    assert path and clusters == ['C1']
    assert FigureEligibilityPolicy().classify({'kind': 'reader_concordance', 'image_path': path, 'labels_readable': True,
          'quantitative_bindings': binding}, citation_complete=False)[0] == 'supplementary'
    rows[0].pop('pair_transition_type_counts')
    assert len(concordance_bindings(state)) == 1  # No rate-to-count reconstruction.


def test_caption_mentions_and_short_fallback_cannot_satisfy_final_report_quality():
    from report_generation.core.reader_authoring import audit_report_output_correctness
    result = audit_report_output_correctness('## Results\n\nAn observed contrast remained.\n\n### Figure 1. Caption only\nA figure.\n',
              {'figures': [{'placement': 'main', 'display_label': 'Figure 1'}]})
    assert result['status'] != 'release_candidate'
    assert result['main_figures_unreferenced_in_prose'] == ['Figure 1']
    assert 'empty_section' in result['section_content_quality']['discussion']


def test_collapsed_annotation_never_routes_mixed_features_by_primary_taxon():
    from rag_enrichment.core.ptm_merger import collapse_ptm_rows_for_enrichment
    from ptm_shared.annotation_species import annotation_scope
    rows = [{'Gene.Name': 'MIXED', 'PTM_Position': 'S1', 'Precursor.Id': 'a', 'Protein.Group': 'rat-1',
             'FASTA_Taxonomy_ID': '10116', 'Condition': '5min', 'PTM_Relative_Log2FC': 3.},
            {'Gene.Name': 'MIXED', 'PTM_Position': 'S1', 'Precursor.Id': 'b', 'Protein.Group': 'human-1',
             'FASTA_Taxonomy_ID': '9606', 'Condition': '5min', 'PTM_Relative_Log2FC': 1.}]
    for ordered in (rows, rows[::-1]):
        work = collapse_ptm_rows_for_enrichment(ordered)[0]
        assert len(work['annotation_feature_provenance']) == 2
        scope = annotation_scope(work, {'organism': 'rat'})
        assert scope['annotation_taxonomy_id'] is None
        assert scope['feature_taxonomy_ids'] == ['10116', '9606']


def test_valid_prose_is_preserved_when_missing_finding_is_recovered():
    from report_generation.core.reader_authoring import build_authoring_packet, render_reader_section_fallback, restore_missing_finding_paragraphs, audit_finding_coverage
    packet = build_authoring_packet(researcher_state())
    valid = render_reader_section_fallback('results', packet).split('\n\n')[0]
    recovered, audit = restore_missing_finding_paragraphs('results', valid, packet)
    assert recovered.startswith(valid) and recovered.count(valid) == 1
    assert audit and audit[0]['status'] == 'deterministic_observation_recovered_review_required'
    assert not audit_finding_coverage({'results': recovered}, packet)['missing_finding_ids']


def test_gene_context_is_allowed_but_exact_site_activation_and_pathway_overreach_are_not():
    from report_generation.core.reader_authoring import build_authoring_packet, validate_and_repair_sections
    state = researcher_state(1)
    for row in state['vector_plot_raw_data']:
        row.update(gene='LARP1', position='T1180')
    packet = build_authoring_packet(state, references=[{'title': 'Synthetic contextual fixture', 'pmid': '1', 'authors': 'Fixture', 'year': '2026'}])
    text = 'Published work described LARP1 in translation control [REF:pmid:1].'
    result, audit = validate_and_repair_sections({'discussion': text}, packet)
    assert result['discussion'] == text
    for claim in ('Our LARP1 T1180 measurement proves a direct mTOR substrate relationship.',
                  'The pathway was significantly enriched with q=0.1756.',
                  'Our AKT1 S122 measurements demonstrate that AKT1 was activated.'):
        result, audit = validate_and_repair_sections({'results': claim}, packet)
        assert claim not in result['results']
        assert any(e['reason_code'] != ['within_contract'] for e in audit['entries'])


def test_two_modifications_are_not_rendered_as_a_single_site_measurement():
    from report_generation.core.reader_authoring import build_authoring_packet
    from report_generation.core.quantitative_claims import validate_quantitative_sentence
    state = researcher_state(1)
    for row in state['vector_plot_raw_data']:
        row.update(modified_sequence='T(UniMod:21)EY(UniMod:21)PEPTIDE')
    packet = build_authoring_packet(state)
    card = next(c for c in packet['reader_cards'] if c.get('trajectory'))
    fid = card['feature_identity']['reader_feature_id']
    assert validate_quantitative_sentence(f'{fid} was a single phosphorylation site measurement.', packet) == ['multimodified_precursor_cannot_be_reduced_to_single_site']


def test_gene_context_retrieval_survives_a_failed_site_search():
    import json
    from report_generation.core.finding_literature import retrieve_finding_literature
    from report_generation.core.measured_feature_cards import build_feature_observation_cards
    class Retriever:
        collection_names = ['fixture']
        def query(self, query, **kwargs):
            if 'opposing results' in query: raise TimeoutError('site search unavailable')
            return [{'document': 'GENE0 participates in transport in rat cells.', 'title': 'Synthetic function source',
                     'pmid': '1', 'source_id': 'fixture-1', 'collection_version': 'fixture-v1'}]
    class Model:
        def generate(self, prompt, **kwargs):
            return json.dumps({'comparisons': [{'source_index': 0, 'quote': 'GENE0 participates in transport in rat cells.',
                'external_finding': 'GENE0 has a transport-related role.', 'reference_scope': 'gene', 'relationship': 'gene_function_context',
                'species': 'rat', 'cell_type': 'rat cells', 'insulin_dose': '', 'time': '', 'readout': '', 'perturbation': ''}]})
    cards = build_feature_observation_cards(researcher_state(1))
    result = retrieve_finding_literature(cards, Retriever(), {'species': 'rat'}, llm=Model())
    record = next(iter(result['records'].values()))
    assert record['status'] == 'context_available' and record['partial_retrieval_failure']
    comparison = result['references'][0]['feature_comparisons'][0]
    assert comparison['reference_scope'] == 'gene' and comparison['measured_relation'] is False


def test_fenced_sentence_json_is_decoded_instead_of_section_fallback():
    import json
    from report_generation.core.quantitative_claims import decode_sentence_draft
    packet = {'reader_cards': [], 'figure_cards': []}
    sentence = {'text': 'The experiment compares recorded precursor contrasts.', 'scope': 'study_rationale',
                'paragraph': 1, 'evidence_ids': [], 'value_tokens': [], 'figure_keys': []}
    fenced = 'Here is the draft:\n```json\n' + json.dumps({'sentences': [sentence]}) + '\n```\n'
    prose, audit = decode_sentence_draft(fenced, packet)
    assert sentence['text'] in prose
    assert all(r.get('retained') for r in audit if 'sentence' in r)


def test_kinase_accuracy_question_is_not_rewritten_as_protein_adjustment():
    from report_generation.core.research_questions import build_question_map
    mapped = build_question_map(['Can we identify the kinase accurately without bias?'], [])
    question = mapped['questions'][0]
    assert 'kinase' in question['normalized_question'].lower()
    assert 'protein adjustment' not in question['normalized_question'].lower()
    assert question['answer_status'] == 'unanswered'


def test_discovery_category_does_not_drop_selected_finding_from_literature_search():
    from report_generation.core.finding_literature import cards_for_selected_findings
    observation = {'category': 'measured_feature_observation', 'feature_identity': {'reader_feature_id': 'PF-1'},
                   'trajectory': [{'ptm_unadjusted_log2fc': 0.2}]}
    discovery = {'category': 'candidate_discovery', 'feature_identity': {'reader_feature_id': 'PF-1'}}
    cards = cards_for_selected_findings([discovery, observation], {'PF-1'})
    assert len(cards) == 1
    assert cards[0]['category'] == 'measured_feature_observation'


def test_literature_target_and_short_complete_sections_are_product_goals_not_empty_defects():
    from report_generation.core.reader_authoring import audit_report_output_correctness
    body = (
        '## Abstract\n\nObserved precursor contrasts were retained for comparison.\n\n'
        '## Introduction\n\nThe study asks which recorded PTM and protein contrasts change together.\n\n'
        '## Methods\n\nIndependent unadjusted, protein, and adjusted contrasts were kept separate.\n\n'
        '## Results\n\nGENEA modified-precursor feature annotated at S10 showed a recorded contrast of +0.20 on the unadjusted PTM contrast.\n\n'
        '## Discussion\n\nFor GENEA modified-precursor feature annotated at S10 the observed contrast remained near the reference level.\n\n'
        '## Conclusion\n\nMeasured contrasts describe the sampled window and do not establish the next validation.\n'
    )
    result = audit_report_output_correctness(body, {'figures': []})
    assert 'literature_coverage_below_product_target' not in result['review_reason_codes']
    assert 'section_content_quality_incomplete' not in result['review_reason_codes']
    assert any('below_section_target' in issues for issues in result['section_content_quality'].values())
